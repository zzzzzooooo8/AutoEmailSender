from __future__ import annotations

import io
import re
import tempfile
from dataclasses import dataclass
from html import escape
from pathlib import Path

from bs4 import BeautifulSoup

from app.services.html_text import html_to_text as convert_html_to_text
from app.services.html_text import normalize_plain_text
from app.models import IdentityProfile, Professor
from app.services.file_storage import extract_text_from_document
from app.services.mail_runtime import text_to_html
from app.services.rich_text import normalize_email_html, text_to_email_html


OUTREACH_GENERATION_MODE_LLM = "llm"
OUTREACH_GENERATION_MODE_TEMPLATE = "template"
TEST_RECIPIENT_NAME = "测试收件人"
SUPPORTED_TEMPLATE_IMPORT_SUFFIXES = {".docx", ".html", ".htm", ".txt", ".md"}
PLACEHOLDER_HELP_TEXT = {
    "name": "导师姓名",
    "email": "导师邮箱",
    "title": "导师职称",
    "university": "导师学校",
    "school": "导师学院",
    "department": "导师院系",
    "research_direction": "导师研究方向",
    "sender_name": "你的发件人姓名",
    "sender_email": "你的发件邮箱",
}
EMAIL_TEMPLATE_FONT_STACK = (
    "'Times New Roman','Songti SC','STSong','SimSun','Noto Serif SC',serif"
)
MISSING_TEMPLATE_SUBJECT_DETAIL = "请先填写默认套磁信主题"
MISSING_TEMPLATE_BODY_TEXT_DETAIL = "请先填写默认套磁信纯文本正文"
MISSING_TEMPLATE_SUBJECT_AND_BODY_TEXT_DETAIL = "请先填写默认套磁信主题和纯文本正文"


@dataclass(slots=True)
class RenderedOutreachTemplate:
    subject: str
    body_text: str
    body_html: str
    placeholders: dict[str, str]


@dataclass(slots=True)
class OutreachTemplateConfig:
    generation_mode: str
    subject_template: str | None
    body_text_template: str | None
    body_html_template: str | None


@dataclass(slots=True)
class ImportedOutreachTemplate:
    subject: str | None
    body_text: str
    body_html: str
    format_name: str


def get_outreach_template_defaults_validation_error(
    subject_template: str | None,
    body_text_template: str | None,
) -> str | None:
    has_subject = bool((subject_template or "").strip())
    has_body_text = bool((body_text_template or "").strip())
    if has_subject and has_body_text:
        return None
    if not has_subject and not has_body_text:
        return MISSING_TEMPLATE_SUBJECT_AND_BODY_TEXT_DETAIL
    if not has_subject:
        return MISSING_TEMPLATE_SUBJECT_DETAIL
    return MISSING_TEMPLATE_BODY_TEXT_DETAIL


def build_template_context(identity: IdentityProfile, professor: Professor) -> dict[str, str]:
    return {
        "name": professor.name or "",
        "email": professor.email or "",
        "title": professor.title or "",
        "university": professor.university or "",
        "school": professor.school or "",
        "department": professor.department or "",
        "research_direction": professor.research_direction or "",
        "sender_name": get_identity_sender_name(identity),
        "sender_email": identity.email_address or "",
    }


def get_identity_sender_name(identity: IdentityProfile) -> str:
    return (
        getattr(identity, "sender_name", None)
        or getattr(identity, "profile_name", None)
        or identity.name
        or ""
    )


def build_test_compose_template_context(identity: IdentityProfile) -> dict[str, str]:
    return {
        "name": TEST_RECIPIENT_NAME,
        "email": identity.email_address or "",
        "title": TEST_RECIPIENT_NAME,
        "university": "测试学校",
        "school": "测试学院",
        "department": "测试院系",
        "research_direction": "测试研究方向",
        "sender_name": get_identity_sender_name(identity),
        "sender_email": identity.email_address or "",
    }


def resolve_outreach_template_config(
    identity: IdentityProfile,
    *,
    generation_mode: str | None = None,
    subject_template: str | None = None,
    body_text_template: str | None = None,
    body_html_template: str | None = None,
) -> OutreachTemplateConfig:
    mode = (generation_mode or identity.outreach_generation_mode or OUTREACH_GENERATION_MODE_LLM).strip()
    if mode not in {OUTREACH_GENERATION_MODE_LLM, OUTREACH_GENERATION_MODE_TEMPLATE}:
        mode = OUTREACH_GENERATION_MODE_LLM

    return OutreachTemplateConfig(
        generation_mode=mode,
        subject_template=(subject_template if subject_template is not None else identity.outreach_template_subject),
        body_text_template=(
            body_text_template
            if body_text_template is not None
            else identity.outreach_template_body_text
        ),
        body_html_template=(
            body_html_template
            if body_html_template is not None
            else identity.outreach_template_body_html
        ),
    )


def render_outreach_template(
    identity: IdentityProfile,
    professor: Professor,
    *,
    subject_template: str | None = None,
    body_text_template: str | None = None,
    body_html_template: str | None = None,
) -> RenderedOutreachTemplate:
    context = build_template_context(identity, professor)
    detail = get_outreach_template_defaults_validation_error(
        subject_template,
        body_text_template,
    )
    if detail:
        raise ValueError(detail)
    normalized_subject_template = (subject_template or "").strip()
    normalized_body_text_template = (body_text_template or "").strip()
    normalized_body_html_template = (body_html_template or "").strip()

    subject = render_template_string(normalized_subject_template, context).strip()
    if not subject:
        raise ValueError("固定模板渲染后缺少邮件主题")

    rendered_body_text = render_template_string(normalized_body_text_template, context).strip()
    rendered_body_html = render_template_string(normalized_body_html_template, context).strip()

    if not rendered_body_text:
        raise ValueError("固定模板渲染后缺少可发送正文")

    if rendered_body_html:
        normalized_html = normalize_html_template(rendered_body_html)
        return RenderedOutreachTemplate(
            subject=subject,
            body_text=rendered_body_text,
            body_html=normalized_html,
            placeholders=context,
        )

    return RenderedOutreachTemplate(
        subject=subject,
        body_text=rendered_body_text,
        body_html=text_to_html(rendered_body_text),
        placeholders=context,
    )


def render_identity_outreach_template(
    identity: IdentityProfile,
    professor: Professor,
) -> RenderedOutreachTemplate:
    return render_outreach_template(
        identity,
        professor,
        subject_template=identity.outreach_template_subject,
        body_text_template=identity.outreach_template_body_text,
        body_html_template=identity.outreach_template_body_html,
    )


def render_template_string(template: str, context: dict[str, str]) -> str:
    def replace(match: re.Match[str]) -> str:
        key = match.group(1).strip()
        return context.get(key, "")

    return re.sub(r"\{\{\s*([a-zA-Z0-9_]+)\s*\}\}", replace, template)


def render_template_with_context(value: str | None, context: dict[str, str]) -> str:
    return render_template_string(value or "", context)


def import_outreach_template_file(file_name: str, content: bytes) -> ImportedOutreachTemplate:
    suffix = Path(file_name).suffix.lower()
    if suffix not in SUPPORTED_TEMPLATE_IMPORT_SUFFIXES:
        raise ValueError("模板文件暂只支持 .docx、.html、.htm、.txt、.md")

    if suffix == ".docx":
        body_html = _convert_docx_template_to_html(content)
        body_text = _extract_docx_template_to_text(content).strip() or html_to_text(body_html)
        if not body_text:
            raise ValueError("模板文件里没有可用正文")
        return ImportedOutreachTemplate(
            subject=None,
            body_text=body_text,
            body_html=body_html,
            format_name=suffix.lstrip("."),
        )

    if suffix in {".html", ".htm"}:
        html_content = _decode_text_file(content).strip()
        if not html_content:
            raise ValueError("模板文件内容为空")
        rendered = normalize_email_html(html_content)
        return ImportedOutreachTemplate(
            subject=None,
            body_text=rendered.text,
            body_html=rendered.html,
            format_name=suffix.lstrip("."),
        )

    if suffix in {".txt", ".md"}:
        body_text = _decode_text_file(content).strip()
    else:
        body_text = _extract_text_from_bytes(file_name, content).strip()

    if not body_text:
        raise ValueError("模板文件里没有可用正文")

    rendered = text_to_email_html(body_text)
    return ImportedOutreachTemplate(
        subject=None,
        body_text=rendered.text,
        body_html=rendered.html,
        format_name=suffix.lstrip("."),
    )


def normalize_html_template(value: str) -> str:
    text = value.strip()
    if not text:
        raise ValueError("HTML 模板内容为空")
    if "<" not in text or ">" not in text:
        return text_to_html(text)

    soup = BeautifulSoup(text, "html.parser")
    if not soup.get_text(" ", strip=True):
        raise ValueError("HTML 模板缺少可见正文")
    return str(soup)


def html_to_text(value: str) -> str:
    return convert_html_to_text(value)


def _decode_text_file(content: bytes) -> str:
    for encoding in ("utf-8-sig", "utf-8", "gb18030"):
        try:
            return content.decode(encoding)
        except UnicodeDecodeError:
            continue
    return content.decode("utf-8", errors="ignore")


def _extract_text_from_bytes(file_name: str, content: bytes) -> str:
    suffix = Path(file_name).suffix.lower()
    with tempfile.NamedTemporaryFile(delete=False, suffix=suffix) as temp_file:
        temp_file.write(content)
        temp_path = Path(temp_file.name)

    try:
        return extract_text_from_document(temp_path.as_posix()) or ""
    finally:
        if temp_path.exists():
            temp_path.unlink()


def _convert_docx_template_to_html(content: bytes) -> str:
    formatted_html = _convert_docx_template_to_formatted_html(content)
    if formatted_html:
        return normalize_html_template(formatted_html)

    try:
        import mammoth
    except ImportError as exc:
        raise ValueError("当前环境缺少 .docx 模板解析依赖 mammoth") from exc

    try:
        result = mammoth.convert_to_html(io.BytesIO(content))
    except Exception as exc:
        raise ValueError("解析 .docx 模板失败") from exc

    html_content = (result.value or "").strip()
    if not html_content:
        raise ValueError("模板文件里没有可用正文")
    return normalize_html_template(_decorate_docx_email_html(html_content))


def _convert_docx_template_to_formatted_html(content: bytes) -> str:
    try:
        from docx import Document
        from docx.oxml.ns import qn
        from docx.table import Table
        from docx.text.paragraph import Paragraph
    except ImportError as exc:
        raise ValueError("当前环境缺少 .docx 模板解析依赖 python-docx") from exc

    try:
        document = Document(io.BytesIO(content))
    except Exception as exc:
        raise ValueError("解析 .docx 模板失败") from exc

    html_parts: list[str] = []
    visible_paragraph_count = 0
    for child in document.element.body.iterchildren():
        if child.tag == qn("w:p"):
            paragraph = Paragraph(child, document)
            paragraph_html = _render_docx_paragraph(paragraph, visible_paragraph_count)
            if paragraph_html:
                visible_paragraph_count += 1
                html_parts.append(paragraph_html)
        elif child.tag == qn("w:tbl"):
            table_html = _render_docx_table(Table(child, document))
            if table_html:
                html_parts.append(table_html)

    if not html_parts:
        return ""

    container = BeautifulSoup("", "html.parser").new_tag(
        "div",
        style=(
            f"font-family:{EMAIL_TEMPLATE_FONT_STACK};"
            "font-size:12pt;"
            "line-height:1.5;"
            "color:#000000;"
            "max-width:100%;"
        ),
    )
    container.append(BeautifulSoup("".join(html_parts), "html.parser"))
    return str(container)


def _render_docx_paragraph(paragraph, visible_index: int) -> str:
    body_html = _render_docx_runs(paragraph.runs)
    text = paragraph.text.strip()
    if not text and not body_html.strip():
        return ""

    tag_name = _get_docx_paragraph_tag_name(paragraph)
    soup = BeautifulSoup("", "html.parser")
    tag = soup.new_tag(tag_name)
    style = _build_docx_paragraph_style(paragraph, tag_name, visible_index)
    if style:
        tag["style"] = style
    tag.append(BeautifulSoup(body_html or escape(text), "html.parser"))
    return str(tag)


def _render_docx_runs(runs) -> str:
    parts: list[str] = []
    for run in runs:
        if not run.text:
            continue
        text = escape(run.text).replace("\n", "<br>")
        style = _build_docx_run_style(run)
        if style:
            text = f'<span style="{style}">{text}</span>'
        if run.bold:
            text = f"<strong>{text}</strong>"
        if run.italic:
            text = f"<em>{text}</em>"
        if run.underline:
            text = f"<u>{text}</u>"
        parts.append(text)
    return "".join(parts)


def _render_docx_table(table) -> str:
    row_parts: list[str] = []
    for row in table.rows:
        cell_parts: list[str] = []
        for cell in row.cells:
            paragraph_parts = [
                _render_docx_paragraph(paragraph, 0)
                for paragraph in cell.paragraphs
            ]
            content = "".join(part for part in paragraph_parts if part)
            cell_parts.append(
                "<td "
                'style="border:1px solid #666666;padding:8px 10px;'
                'vertical-align:top;font-size:12pt;line-height:1.5;color:#000000;">'
                f"{content}</td>"
            )
        row_parts.append(f"<tr>{''.join(cell_parts)}</tr>")

    if not row_parts:
        return ""
    return (
        '<table style="width:100%;border-collapse:collapse;table-layout:fixed;'
        f'margin:12px 0;"><tbody>{"".join(row_parts)}</tbody></table>'
    )


def _get_docx_paragraph_tag_name(paragraph) -> str:
    style_name = (getattr(paragraph.style, "name", "") or "").lower()
    match = re.search(r"heading\s+([1-6])", style_name)
    if match:
        return f"h{match.group(1)}"
    return "p"


def _build_docx_paragraph_style(paragraph, tag_name: str, visible_index: int) -> str:
    fragments = [
        "margin:0 0 12px 0",
        f"font-family:{_get_docx_paragraph_font_family(paragraph)}",
        f"font-size:{_get_docx_paragraph_font_size(paragraph, tag_name)}",
        f"line-height:{_get_docx_paragraph_line_height(paragraph)}",
        "color:#000000",
        "white-space:pre-wrap",
    ]

    text_align = _get_docx_paragraph_text_align(paragraph)
    if text_align:
        fragments.append(f"text-align:{text_align}")

    indent = _get_docx_first_line_indent(paragraph)
    if tag_name == "p":
        fragments.append(f"text-indent:{indent if indent is not None else '0'}")
    if visible_index == 0 and indent is None:
        fragments.append("text-indent:0")

    if tag_name.startswith("h"):
        fragments.extend(["font-weight:700", "text-indent:0"])

    return ";".join(fragments) + ";"


def _build_docx_run_style(run) -> str:
    fragments: list[str] = []
    font_family = _get_docx_run_font_family(run)
    if font_family:
        fragments.append(f"font-family:{font_family}")
    font_size = _get_docx_run_font_size(run)
    if font_size:
        fragments.append(f"font-size:{font_size}")
    color = _get_docx_run_color(run)
    if color:
        fragments.append(f"color:{color}")
    return ";".join(fragments) + (";" if fragments else "")


def _get_docx_paragraph_font_family(paragraph) -> str:
    for run in paragraph.runs:
        font_family = _get_docx_run_font_family(run)
        if font_family:
            return font_family
    return EMAIL_TEMPLATE_FONT_STACK


def _get_docx_paragraph_font_size(paragraph, tag_name: str) -> str:
    for run in paragraph.runs:
        font_size = _get_docx_run_font_size(run)
        if font_size:
            return font_size
    heading_defaults = {
        "h1": "16pt",
        "h2": "14pt",
        "h3": "13pt",
        "h4": "12pt",
        "h5": "12pt",
        "h6": "12pt",
    }
    return heading_defaults.get(tag_name, "12pt")


def _get_docx_paragraph_line_height(paragraph) -> str:
    line_spacing = paragraph.paragraph_format.line_spacing
    if isinstance(line_spacing, (int, float)):
        return f"{line_spacing:g}"
    return "1.5"


def _get_docx_paragraph_text_align(paragraph) -> str | None:
    alignment = paragraph.alignment
    if alignment is None:
        return None
    name = getattr(alignment, "name", "").lower()
    if "center" in name:
        return "center"
    if "right" in name:
        return "right"
    if "justify" in name or "distribute" in name:
        return "justify"
    if "left" in name:
        return "left"
    return None


def _get_docx_first_line_indent(paragraph) -> str | None:
    first_line_indent = paragraph.paragraph_format.first_line_indent
    if first_line_indent is None:
        return None
    points = _docx_length_to_points(first_line_indent)
    if points is None or points <= 0:
        return "0"
    font_size = _docx_font_size_to_points(_get_docx_paragraph_font_size(paragraph, "p")) or 12.0
    em_value = points / font_size
    if abs(em_value - round(em_value)) < 0.05:
        em_text = str(int(round(em_value)))
    else:
        em_text = f"{em_value:.2f}".rstrip("0").rstrip(".")
    return f"{em_text}em"


def _get_docx_run_font_family(run) -> str | None:
    font_name = run.font.name
    rpr = getattr(run._element, "rPr", None)
    rfonts = getattr(rpr, "rFonts", None) if rpr is not None else None
    east_asia_name = None
    if rfonts is not None:
        try:
            from docx.oxml.ns import qn

            east_asia_name = rfonts.get(qn("w:eastAsia"))
        except Exception:
            east_asia_name = None
    if east_asia_name and font_name and east_asia_name != font_name:
        return f"'{font_name}','{east_asia_name}',{EMAIL_TEMPLATE_FONT_STACK}"
    if east_asia_name:
        return f"'{east_asia_name}',{EMAIL_TEMPLATE_FONT_STACK}"
    if font_name:
        return f"'{font_name}',{EMAIL_TEMPLATE_FONT_STACK}"
    return None


def _get_docx_run_font_size(run) -> str | None:
    if run.font.size is None:
        return None
    points = _docx_length_to_points(run.font.size)
    if points is None:
        return None
    return f"{points:g}pt"


def _get_docx_run_color(run) -> str | None:
    color = run.font.color
    if color is None or color.rgb is None:
        return None
    return f"#{color.rgb}"


def _docx_length_to_points(value) -> float | None:
    if value is None:
        return None
    points = getattr(value, "pt", None)
    if points is not None:
        return float(points)
    try:
        return float(value) / 12700
    except (TypeError, ValueError):
        return None


def _docx_font_size_to_points(value: str) -> float | None:
    match = re.match(r"^([0-9]+(?:\.[0-9]+)?)pt$", value)
    if not match:
        return None
    return float(match.group(1))


def _extract_docx_template_to_text(content: bytes) -> str:
    try:
        import mammoth
    except ImportError as exc:
        raise ValueError("当前环境缺少 .docx 模板解析依赖 mammoth") from exc

    try:
        result = mammoth.extract_raw_text(io.BytesIO(content))
    except Exception as exc:
        raise ValueError("解析 .docx 模板失败") from exc

    return normalize_plain_text(result.value or "")


def _decorate_docx_email_html(value: str) -> str:
    soup = BeautifulSoup(value, "html.parser")
    container = soup.new_tag(
        "div",
        style=(
            f"font-family:{EMAIL_TEMPLATE_FONT_STACK};"
            "font-size:12pt;"
            "line-height:1.5;"
            "color:#000000;"
            "max-width:100%;"
        ),
    )
    for child in list(soup.contents):
        container.append(child.extract())

    first_paragraph = next(
        (tag for tag in container.find_all("p") if tag.get_text(" ", strip=True)),
        None,
    )
    for paragraph in container.find_all("p"):
        text = paragraph.get_text(" ", strip=True)
        is_in_table = paragraph.find_parent("table") is not None
        is_section_heading = bool(re.match(r"^（[一二三四五六七八九十]+）", text))
        indent = "0" if paragraph is first_paragraph or is_in_table or is_section_heading else "2em"
        margin_top = "16px" if is_section_heading else "0"
        _append_inline_style(
            paragraph,
            (
                f"margin:{margin_top} 0 12px 0;"
                "font-size:12pt;"
                "line-height:1.5;"
                "color:#000000;"
                f"text-indent:{indent};"
            ),
        )

    for strong_tag in container.find_all(["strong", "b"]):
        _append_inline_style(strong_tag, "font-weight:700;")

    for table in container.find_all("table"):
        _append_inline_style(
            table,
            (
                "width:100%;"
                "border-collapse:collapse;"
                "table-layout:fixed;"
                "margin:12px 0;"
            ),
        )

    for cell in container.find_all(["td", "th"]):
        _append_inline_style(
            cell,
            (
                "border:1px solid #666666;"
                "padding:8px 10px;"
                "vertical-align:top;"
                "font-size:12pt;"
                "line-height:1.5;"
                "color:#000000;"
            ),
        )

    for list_tag in container.find_all(["ul", "ol"]):
        _append_inline_style(
            list_tag,
            "margin:0 0 12px 1.5em;padding:0;line-height:1.5;",
        )

    for item in container.find_all("li"):
        _append_inline_style(item, "margin:0 0 6px 0;")

    return str(container)


def _append_inline_style(tag, style_fragment: str) -> None:
    existing_style = (tag.get("style") or "").strip()
    if existing_style and not existing_style.endswith(";"):
        existing_style = f"{existing_style};"
    tag["style"] = f"{existing_style}{style_fragment}"
